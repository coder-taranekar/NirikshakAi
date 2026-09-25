"""
Report Generation Service — PDF, Excel, and DOCX compliance reports.

Three formats:
  pdf   — WeasyPrint renders the Jinja2 HTML template to a print-ready PDF
  xlsx  — openpyxl builds a multi-sheet workbook with conditional formatting
  docx  — python-docx builds an editable Word document with styled tables

Entry point: generate_report(inspection, fmt) → (bytes, content_type, filename)
Called by GET /inspections/{id}/report in the inspections router.
"""

import io
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from jinja2 import Environment, FileSystemLoader, select_autoescape

from app.models.inspection import Inspection

_TEMPLATES_DIR = Path(__file__).parent.parent / "templates"

# ── Jinja2 Environment ─────────────────────────────────────────────────────────

def _make_jinja_env() -> Environment:
    env = Environment(
        loader=FileSystemLoader(str(_TEMPLATES_DIR)),
        autoescape=select_autoescape(["html"]),
    )

    def format_inr(value: int) -> str:
        """Format an integer as Indian Rupee string with commas: 100000 → 1,00,000"""
        s = str(abs(value))
        if len(s) <= 3:
            return s
        result = s[-3:]
        s = s[:-3]
        while s:
            result = s[-2:] + "," + result
            s = s[:-2]
        return ("-" if value < 0 else "") + result

    env.filters["format_inr"] = format_inr
    return env


_jinja_env = _make_jinja_env()


# ── Template Context Builder ───────────────────────────────────────────────────

def _build_context(inspection: Inspection) -> dict[str, Any]:
    """Build the template rendering context from an Inspection ORM object."""
    cr = inspection.compliance_result or {}
    ocr = inspection.ocr_result or {}
    advanced = cr.get("advanced_checks", {})
    penalty = cr.get("penalty", {})
    declarations = cr.get("declarations", [])

    # Product & manufacturer info
    product = inspection.product
    product_name = product.name if product else "Unknown"
    generic_name = product.generic_name if product else "—"
    brand = product.brand if product else "—"
    barcode = product.barcode if product else "—"
    schedule_ii = "Yes" if (product and product.is_schedule_ii_commodity) else "No"

    manufacturer = "—"
    if product and product.manufacturer:
        manufacturer = product.manufacturer.name

    inspector_name = "—"
    if inspection.inspector:
        inspector_name = inspection.inspector.name

    status = inspection.status.value if inspection.status else "pending"
    status_labels = {
        "compliant": "Compliant",
        "partial": "Partially Compliant",
        "non_compliant": "Non-Compliant",
        "pending": "Pending",
    }

    return {
        "inspection_id": str(inspection.id),
        "generated_at": datetime.now(tz=timezone.utc).strftime("%d %b %Y %H:%M UTC"),
        "inspection_date": inspection.created_at.strftime("%d %b %Y %H:%M UTC")
        if inspection.created_at else "—",
        "inspector_name": inspector_name,
        "source": inspection.source.value if inspection.source else "—",
        "source_url": inspection.source_url or "",
        "state": inspection.state or "—",
        "district": inspection.district or "—",
        "score": round(inspection.score or 0, 1),
        "status": status,
        "status_label": status_labels.get(status, status),
        "hard_fail": inspection.hard_fail_triggered,
        "product_name": product_name,
        "generic_name": generic_name,
        "brand": brand,
        "manufacturer": manufacturer,
        "barcode": barcode,
        "schedule_ii": schedule_ii,
        "declarations": declarations,
        "advanced": advanced,
        "penalty": penalty,
        "inspector_notes": inspection.inspector_notes or "",
    }


# ── PDF Generator ──────────────────────────────────────────────────────────────

def _generate_pdf(inspection: Inspection) -> bytes:
    """Render the HTML template and convert to PDF via WeasyPrint."""
    from weasyprint import HTML  # type: ignore

    context = _build_context(inspection)
    template = _jinja_env.get_template("inspection_report.html")
    html_content = template.render(**context)
    pdf_bytes = HTML(string=html_content, base_url=str(_TEMPLATES_DIR)).write_pdf()
    return pdf_bytes


# ── Excel Generator ────────────────────────────────────────────────────────────

_STATUS_FILL_MAP = {
    "pass":  ("D1FAE5", "065F46"),  # green bg, dark green text
    "fail":  ("FEE2E2", "991B1B"),  # red bg, dark red text
}

_HEADER_FILL = "1A1A2E"  # dark navy
_HEADER_FONT = "FFFFFF"  # white


def _generate_xlsx(inspection: Inspection) -> bytes:
    """Build a multi-sheet Excel workbook with conditional formatting."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # ── Sheet 1: Summary ────────────────────────────────────────────
    ws_summary = wb.active
    ws_summary.title = "Summary"

    ctx = _build_context(inspection)

    def hdr_cell(ws, row, col, value):
        c = ws.cell(row=row, column=col, value=value)
        c.font = Font(bold=True, color=_HEADER_FONT)
        c.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        c.alignment = Alignment(vertical="center")
        return c

    def data_cell(ws, row, col, value, bold=False, bg=None, fg=None):
        c = ws.cell(row=row, column=col, value=value)
        if bold:
            c.font = Font(bold=True, color=fg or "000000")
        if bg:
            c.fill = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        return c

    # Title row
    ws_summary.merge_cells("A1:D1")
    title_cell = ws_summary["A1"]
    title_cell.value = f"LabelGuard Compliance Report — {ctx['inspection_id']}"
    title_cell.font = Font(bold=True, size=14, color=_HEADER_FILL)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 24

    # Meta info
    meta_rows = [
        ("Generated", ctx["generated_at"]),
        ("Inspector", ctx["inspector_name"]),
        ("Inspection Date", ctx["inspection_date"]),
        ("Source", ctx["source"]),
        ("State / District", f"{ctx['state']} / {ctx['district']}"),
        ("", ""),
        ("Compliance Score", f"{ctx['score']} / 100"),
        ("Status", ctx["status_label"]),
        ("Hard Fail Triggered", "Yes" if ctx["hard_fail"] else "No"),
        ("", ""),
        ("Product Name", ctx["product_name"]),
        ("Generic Name", ctx["generic_name"]),
        ("Brand", ctx["brand"]),
        ("Manufacturer", ctx["manufacturer"]),
        ("Barcode", ctx["barcode"]),
        ("Schedule II", ctx["schedule_ii"]),
    ]

    for i, (label, value) in enumerate(meta_rows, start=3):
        ws_summary.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws_summary.cell(row=i, column=2, value=value)

    ws_summary.column_dimensions["A"].width = 24
    ws_summary.column_dimensions["B"].width = 50

    # Penalty section
    penalty = ctx["penalty"]
    pen_row = len(meta_rows) + 4
    hdr_cell(ws_summary, pen_row, 1, "Penalty Assessment")
    hdr_cell(ws_summary, pen_row, 2, "Details")
    ws_summary.cell(row=pen_row + 1, column=1, value="Section").font = Font(bold=True)
    ws_summary.cell(row=pen_row + 1, column=2, value=penalty.get("applicable_section", "—"))
    ws_summary.cell(row=pen_row + 2, column=1, value="Offence Tier").font = Font(bold=True)
    ws_summary.cell(row=pen_row + 2, column=2, value=(penalty.get("tier") or "None").upper())
    ws_summary.cell(row=pen_row + 3, column=1, value="Estimated Penalty").font = Font(bold=True)
    ws_summary.cell(
        row=pen_row + 3, column=2,
        value=f"₹{penalty.get('estimated_min', 0):,} – ₹{penalty.get('estimated_max', 0):,}"
    )
    ws_summary.cell(row=pen_row + 4, column=1, value="Note").font = Font(bold=True)
    ws_summary.cell(row=pen_row + 4, column=2, value=penalty.get("note", "—"))

    # ── Sheet 2: Declaration Checks ─────────────────────────────────
    ws_decl = wb.create_sheet("Declaration Checks")
    headers = ["Field", "Label", "Result", "Rule", "Extracted Value", "Violation Description", "Weight"]
    for col, h in enumerate(headers, start=1):
        c = hdr_cell(ws_decl, 1, col, h)

    for row_idx, d in enumerate(ctx["declarations"], start=2):
        status = d.get("status", "")
        bg, fg = _STATUS_FILL_MAP.get(status, ("FFFFFF", "000000"))
        ws_decl.cell(row=row_idx, column=1, value=d.get("field", ""))
        ws_decl.cell(row=row_idx, column=2, value=d.get("label", ""))
        result_cell = ws_decl.cell(
            row=row_idx, column=3,
            value="✓ PASS" if status == "pass" else "✗ FAIL"
        )
        result_cell.fill = PatternFill("solid", fgColor=bg)
        result_cell.font = Font(bold=True, color=fg)
        ws_decl.cell(row=row_idx, column=4, value=d.get("rule_reference", ""))
        ws_decl.cell(row=row_idx, column=5, value=d.get("extracted_value") or "—")
        ws_decl.cell(row=row_idx, column=6, value=d.get("violation_description", "") if status == "fail" else "—")
        ws_decl.cell(row=row_idx, column=7, value=d.get("weight", 0))

    col_widths = [18, 32, 10, 12, 30, 40, 8]
    for i, w in enumerate(col_widths, start=1):
        ws_decl.column_dimensions[get_column_letter(i)].width = w

    # ── Sheet 3: Advanced Checks ────────────────────────────────────
    ws_adv = wb.create_sheet("Advanced Checks")
    advanced = ctx["advanced"]

    adv_rows = [
        ("Font Size Compliant", "Yes" if advanced.get("font_size_compliant") else "No",
         "D1FAE5" if advanced.get("font_size_compliant") else "FEE2E2"),
        ("Font Size Violations", str(advanced.get("font_size_failed", 0)), None),
        ("Whitespace Compliant", "Yes" if advanced.get("whitespace_compliant") else "No",
         "D1FAE5" if advanced.get("whitespace_compliant") else "FEE2E2"),
        ("MRP Sticker Fraud", "Detected" if advanced.get("sticker_fraud_detected") else "Clear",
         "FEE2E2" if advanced.get("sticker_fraud_detected") else "D1FAE5"),
        ("MRP Values Found", ", ".join(advanced.get("mrp_values_found", [])) or "—", None),
        ("Pack Size Violation", "Yes" if advanced.get("pack_size_violation") else "No / N/A",
         "FEE2E2" if advanced.get("pack_size_violation") else "D1FAE5"),
        ("Pack Size Note", advanced.get("pack_size_note", "—"), None),
        ("Hindi/Devanagari", "Present" if advanced.get("devanagari_detected") else "Not Detected",
         "D1FAE5" if advanced.get("devanagari_detected") else "FFF7ED"),
        ("Multilingual Note", advanced.get("multilingual_note", "—"), None),
    ]

    hdr_cell(ws_adv, 1, 1, "Check")
    hdr_cell(ws_adv, 1, 2, "Result")
    for i, (check, result, bg) in enumerate(adv_rows, start=2):
        ws_adv.cell(row=i, column=1, value=check).font = Font(bold=True)
        c = ws_adv.cell(row=i, column=2, value=result)
        if bg:
            c.fill = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(wrap_text=True)

    ws_adv.column_dimensions["A"].width = 28
    ws_adv.column_dimensions["B"].width = 50

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── DOCX Generator ─────────────────────────────────────────────────────────────

def _generate_docx(inspection: Inspection) -> bytes:
    """Build an editable Word document compliance report."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    ctx = _build_context(inspection)

    # ── Styles ────────────────────────────────────────────────────────
    def add_heading(text: str, level: int = 1):
        p = doc.add_heading(text, level=level)
        if level == 1:
            p.runs[0].font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        return p

    def add_kv(label: str, value: str):
        p = doc.add_paragraph()
        r = p.add_run(f"{label}: ")
        r.bold = True
        p.add_run(value)
        p.paragraph_format.space_after = Pt(0)

    # ── Title ─────────────────────────────────────────────────────────
    title = doc.add_heading("LabelGuard Compliance Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "Legal Metrology (Packaged Commodities) Rules, 2011\n"
        "Department of Consumer Affairs, Government of India"
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # ── Meta ──────────────────────────────────────────────────────────
    add_heading("Inspection Details", 1)
    add_kv("Report ID", ctx["inspection_id"])
    add_kv("Generated", ctx["generated_at"])
    add_kv("Inspector", ctx["inspector_name"])
    add_kv("Inspection Date", ctx["inspection_date"])
    add_kv("Source", ctx["source"])
    if ctx["source_url"]:
        add_kv("URL", ctx["source_url"])
    add_kv("State / District", f"{ctx['state']} / {ctx['district']}")
    add_kv("Compliance Score", f"{ctx['score']} / 100")
    add_kv("Status", ctx["status_label"])
    add_kv("Hard Fail", "Yes ⚠" if ctx["hard_fail"] else "No")

    doc.add_paragraph()
    add_heading("Product Information", 1)
    add_kv("Product Name", ctx["product_name"])
    add_kv("Generic Name", ctx["generic_name"])
    add_kv("Brand", ctx["brand"])
    add_kv("Manufacturer", ctx["manufacturer"])
    add_kv("Barcode", ctx["barcode"])
    add_kv("Schedule II Commodity", ctx["schedule_ii"])

    # ── Declaration Checks Table ──────────────────────────────────────
    doc.add_paragraph()
    add_heading("Mandatory Declaration Checks (Rule 6)", 1)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(["Declaration", "Result", "Rule", "Extracted Value", "Violation"]):
        hdr_cells[i].text = h
        for run in hdr_cells[i].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # Navy background
        tc = hdr_cells[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1A1A2E")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)

    for d in ctx["declarations"]:
        row_cells = table.add_row().cells
        row_cells[0].text = d.get("label", "")
        result_text = "✓ PASS" if d["status"] == "pass" else "✗ FAIL"
        row_cells[1].text = result_text
        row_cells[2].text = d.get("rule_reference", "")
        row_cells[3].text = d.get("extracted_value") or "—"
        row_cells[4].text = d.get("violation_description", "") if d["status"] == "fail" else "—"

        # Colour the result cell
        fill_color = "D1FAE5" if d["status"] == "pass" else "FEE2E2"
        tc = row_cells[1]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill_color)
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)

    # ── Penalty ────────────────────────────────────────────────────────
    doc.add_paragraph()
    add_heading("Penalty Assessment", 1)
    penalty = ctx["penalty"]
    add_kv("Applicable Section", penalty.get("applicable_section", "—"))
    add_kv("Offence Tier", (penalty.get("tier") or "None").upper())
    add_kv(
        "Estimated Penalty",
        f"₹{penalty.get('estimated_min', 0):,} – ₹{penalty.get('estimated_max', 0):,}"
    )
    add_kv("Note", penalty.get("note", "—"))

    # ── Inspector Notes (editable section) ────────────────────────────
    doc.add_paragraph()
    add_heading("Inspector Notes", 1)
    if ctx["inspector_notes"]:
        doc.add_paragraph(ctx["inspector_notes"])
    else:
        p = doc.add_paragraph("[Add your notes here]")
        p.runs[0].font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
        p.runs[0].italic = True

    # ── Footer ─────────────────────────────────────────────────────────
    doc.add_paragraph()
    footer_p = doc.add_paragraph(
        f"Report ID: {ctx['inspection_id']} | Generated: {ctx['generated_at']}\n"
        "LabelGuard — Automated Compliance Checking System"
    )
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer_p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Main Entry Point ───────────────────────────────────────────────────────────

def generate_report(
    inspection: Inspection,
    fmt: str,
) -> tuple[bytes, str, str]:
    """
    Generate a compliance report in the requested format.

    Args:
        inspection: ORM object (with relationships loaded or lazy-loadable)
        fmt: "pdf" | "xlsx" | "docx"

    Returns:
        (report_bytes, content_type, filename)
    """
    short_id = str(inspection.id)[:8]
    date_str = (inspection.created_at or datetime.now(tz=timezone.utc)).strftime("%Y%m%d")
    base_name = f"labelguard_report_{short_id}_{date_str}"

    if fmt == "pdf":
        data = _generate_pdf(inspection)
        return data, "application/pdf", f"{base_name}.pdf"

    if fmt == "xlsx":
        data = _generate_xlsx(inspection)
        return (
            data,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{base_name}.xlsx",
        )

    if fmt == "docx":
        data = _generate_docx(inspection)
        return (
            data,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{base_name}.docx",
        )

    raise ValueError(f"Unsupported report format: {fmt}")
